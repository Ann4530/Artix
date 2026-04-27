#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import sys
import os
from pathlib import Path

# Try to import required libraries
try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Installing required libraries...")
    os.system("pip install python-docx -q")
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

def read_markdown(file_path):
    """Read markdown file"""
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()

def markdown_to_docx(markdown_path, docx_path):
    """Convert markdown to Word document"""
    # Read markdown content
    with open(markdown_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    # Create a new Document
    doc = Document()
    
    # Add document properties
    doc.core_properties.title = "ARTIX - Use Cases"
    doc.core_properties.author = "Artix Team"
    
    # Process markdown lines
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')
        
        # Handle headings
        if line.startswith('# '):
            heading = line[2:]
            p = doc.add_heading(heading, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            heading = line[3:]
            doc.add_heading(heading, level=2)
        elif line.startswith('### '):
            heading = line[4:]
            doc.add_heading(heading, level=3)
        elif line.startswith('#### '):
            heading = line[5:]
            doc.add_heading(heading, level=4)
        
        # Handle bullet lists
        elif line.strip().startswith('- '):
            bullet_text = line.strip()[2:]
            # Check for bold/italic in markdown
            bullet_text = process_inline_formatting(bullet_text)
            doc.add_paragraph(bullet_text, style='List Bullet')
        
        # Handle table
        elif line.strip().startswith('|'):
            # Parse table
            table_rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [cell.strip() for cell in lines[i].split('|')[1:-1]]
                table_rows.append(row)
                i += 1
            i -= 1
            
            if table_rows:
                # Create table
                table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                table.style = 'Light Grid Accent 1'
                
                for row_idx, row_data in enumerate(table_rows):
                    for col_idx, cell_data in enumerate(row_data):
                        cell = table.rows[row_idx].cells[col_idx]
                        cell.text = cell_data
                        # Make header row bold
                        if row_idx == 0:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.font.bold = True
        
        # Handle empty lines
        elif not line.strip():
            doc.add_paragraph()
        
        # Handle regular text
        elif line.strip():
            # Process inline formatting
            text = process_inline_formatting(line.strip())
            doc.add_paragraph(text)
        
        i += 1
    
    # Save document
    doc.save(docx_path)
    print(f"✓ Successfully created: {docx_path}")

def process_inline_formatting(text):
    """Process markdown inline formatting (bold, italic, code)"""
    # Handle bold **text**
    import re
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    # Handle italic *text*
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    # Handle code `text`
    text = re.sub(r'`(.*?)`', r'\1', text)
    return text

if __name__ == '__main__':
    md_file = r'c:\Users\Phuong Anh\Downloads\Project\Artix\SWP391-Artix\uc_specs\ARTIX_USE_CASES_SUMMARY.md'
    docx_file = r'c:\Users\Phuong Anh\Downloads\Project\Artix\SWP391-Artix\uc_specs\ARTIX_USE_CASES_SUMMARY.docx'
    
    if os.path.exists(md_file):
        markdown_to_docx(md_file, docx_file)
    else:
        print(f"File not found: {md_file}")
